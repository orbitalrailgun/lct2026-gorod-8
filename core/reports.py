"""Отчёты в XLSX.

ТЗ требует формирования аналитических отчётов для руководства. Отчёт строится
по оперативному контуру, а не по сырым данным: руководителя интересует, что
система предсказала, что диспетчер с этим сделал и сколько срабатываний удалось
снять как несущественные.

Отдельный лист посвящён наблюдаемости. Это не статистика ради статистики:
доля вердиктов, которые невозможно подтвердить приборами, — прямое основание
для заявки на дооснащение объектов, и заказчик такой цифры сегодня не имеет.
"""

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_sheet(ws, headers, rows, widths=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def build_report(con):
    """Сводный отчёт: вердикты, решения, наблюдаемость, сводка."""
    wb = Workbook()

    with con.cursor() as cur:
        cur.execute("""
            SELECT v.id, v.created_at, v.scenario, v.object_name, v.picket,
                   v.probability, v.horizon_hours, v.status,
                   d.action, r.title AS reason_title, u.full_name, d.decided_at,
                   v.card->>'title'            AS card_title,
                   v.card->>'incubation_days'  AS incubation,
                   v.card->'blind_spots'       AS blind
            FROM verdict v
            LEFT JOIN decision d ON d.verdict_id = v.id
            LEFT JOIN reason_ref r ON r.id = d.reason_id
            LEFT JOIN app_user u ON u.id = d.user_id
            ORDER BY v.probability DESC
        """)
        cols = [c.name for c in cur.description]
        data = [dict(zip(cols, r)) for r in cur.fetchall()]

    # --- лист «Вердикты»
    ws = wb.active
    ws.title = "Вердикты"
    _write_sheet(
        ws,
        ["№", "Сформирован", "Сценарий", "Объект", "ПК", "Риск",
         "Горизонт, ч", "Статус", "Что произошло", "Вызревание, сут",
         "Решение", "Причина", "Диспетчер"],
        [[r["id"],
          r["created_at"].strftime("%d.%m.%Y %H:%M") if r["created_at"] else "",
          r["scenario"], r["object_name"], r["picket"],
          round(float(r["probability"]), 3), r["horizon_hours"], r["status"],
          r["card_title"], r["incubation"], r["action"], r["reason_title"],
          r["full_name"]]
         for r in data],
        widths=[6, 17, 15, 24, 7, 8, 11, 12, 52, 15, 14, 22, 22],
    )

    # --- лист «Сводка»
    ws = wb.create_sheet("Сводка")
    scen = Counter(r["scenario"] for r in data)
    status = Counter(r["status"] for r in data)
    closed = [r for r in data if r["status"] == "закрыт"]
    false_alarm = [r for r in closed if r["reason_title"] in
                   ("Ложное срабатывание", "Профилактика",
                    "Плановая проверка", "Технологические испытания")]

    rows = [["Всего вердиктов", len(data)]]
    rows += [[f"  из них {k}", v] for k, v in scen.most_common()]
    rows += [["", ""]]
    rows += [[f"Статус: {k}", v] for k, v in status.most_common()]
    rows += [["", ""],
             ["Отработано диспетчером", len(closed)],
             ["Признано несущественными", len(false_alarm)]]
    if closed:
        rows.append(["Доля снятых как несущественные",
                     f"{len(false_alarm) / len(closed):.0%}"])
    _write_sheet(ws, ["Показатель", "Значение"], rows, widths=[46, 16])

    # --- лист «Наблюдаемость»
    ws = wb.create_sheet("Наблюдаемость")
    blind = [r for r in data if r["blind"]]
    by_object = Counter(r["object_name"] for r in blind)
    total_by_object = Counter(r["object_name"] for r in data)
    rows = []
    for obj, n_blind in by_object.most_common():
        total = total_by_object[obj]
        rows.append([obj, total, n_blind, f"{n_blind / total:.0%}"])
    _write_sheet(
        ws,
        ["Объект", "Вердиктов", "С ограниченной наблюдаемостью", "Доля"],
        rows, widths=[30, 13, 32, 10])
    ws.append([])
    ws.append(["Ограниченная наблюдаемость означает, что подтвердить вердикт"])
    ws.append(["приборами невозможно: в точке нет датчика нужного типа либо"])
    ws.append(["соседние каналы молчат. Это основание для дооснащения объекта."])

    return wb


def to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
